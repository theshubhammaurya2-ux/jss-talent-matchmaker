"""Extract selectable text from PDF or DOCX job descriptions."""
from io import BytesIO
from pathlib import Path
import pymupdf as fitz
from docx import Document

MAX_FILE_BYTES = 10 * 1024 * 1024
MIN_TEXT_CHARS = 80

class JDDocumentError(ValueError):
    """File is unsupported, unreadable, or lacks selectable text."""


def extract_text(filename: str, file_bytes: bytes) -> str:
    if not file_bytes:
        raise JDDocumentError("The uploaded file is empty.")
    if len(file_bytes) > MAX_FILE_BYTES:
        raise JDDocumentError("Upload a PDF or DOCX smaller than 10 MB.")
    ext = Path(filename).suffix.lower()
    if ext not in {".pdf", ".docx"}:
        raise JDDocumentError("Only PDF and DOCX files are supported.")
    try:
        if ext == ".pdf":
            if not file_bytes.startswith(b"%PDF-"):
                raise JDDocumentError("The file does not have a valid PDF header.")
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                text = "\n".join(page.get_text(sort=True) for page in doc)
        else:
            if not file_bytes.startswith(b"PK"):
                raise JDDocumentError("The file does not appear to be a valid DOCX.")
            doc = Document(BytesIO(file_bytes))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
    except JDDocumentError:
        raise
    except Exception as exc:
        raise JDDocumentError(f"Could not read document: {type(exc).__name__}") from exc

    cleaned = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(cleaned) < MIN_TEXT_CHARS:
        raise JDDocumentError(
            "Almost no selectable text found. A scanned/image-only PDF needs OCR."
        )
    return cleaned
